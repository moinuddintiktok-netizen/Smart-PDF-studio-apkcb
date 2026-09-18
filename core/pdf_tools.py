from pathlib import Path


def pdf_to_word(src, out):
    import fitz
    from docx import Document
    doc = Document()
    pdf = fitz.open(src)
    for i, page in enumerate(pdf):
        text = page.get_text('text')
        if text.strip():
            if i: doc.add_page_break()
            for line in text.splitlines():
                doc.add_paragraph(line)
    doc.save(out)


def word_to_pdf(src, out):
    # Mobile-safe basic DOCX text renderer. Complex Word layouts need a full office engine.
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    docx = Document(src)
    styles = getSampleStyleSheet()
    story=[]
    for p in docx.paragraphs:
        text = p.text.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
        if text.strip(): story.append(Paragraph(text, styles['BodyText']))
        else: story.append(Spacer(1,8))
    if not story: story.append(Paragraph('Empty document', styles['BodyText']))
    SimpleDocTemplate(out, pagesize=A4, rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36).build(story)


def images_to_pdf(images, out):
    from PIL import Image
    ims=[]
    for f in images:
        im=Image.open(f).convert('RGB')
        ims.append(im)
    if not ims: raise ValueError('No images selected')
    ims[0].save(out, save_all=True, append_images=ims[1:])


def merge_pdfs(files, out):
    from pypdf import PdfWriter, PdfReader
    writer=PdfWriter()
    for f in files:
        for page in PdfReader(f).pages: writer.add_page(page)
    with open(out,'wb') as fh: writer.write(fh)


def split_pdf(src, out, start_page, end_page):
    from pypdf import PdfWriter, PdfReader
    reader=PdfReader(src); writer=PdfWriter()
    start=max(1,int(start_page)); end=min(len(reader.pages),int(end_page))
    if start>end: raise ValueError('Invalid page range')
    for i in range(start-1,end): writer.add_page(reader.pages[i])
    with open(out,'wb') as fh: writer.write(fh)


def protect_pdf(src,out,password):
    from pypdf import PdfReader, PdfWriter
    reader=PdfReader(src); writer=PdfWriter()
    for p in reader.pages: writer.add_page(p)
    writer.encrypt(password)
    with open(out,'wb') as fh: writer.write(fh)


def unlock_pdf(src,out,password):
    from pypdf import PdfReader, PdfWriter
    reader=PdfReader(src)
    if reader.is_encrypted and not reader.decrypt(password): raise ValueError('Incorrect password')
    writer=PdfWriter()
    for p in reader.pages: writer.add_page(p)
    with open(out,'wb') as fh: writer.write(fh)
